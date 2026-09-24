#!/usr/bin/env python3
"""Turn a scraped TikTok comments CSV (from tiktok_comments.py) into:
  1) an Excel workbook with tabs: Overview, Top Liked, Comments, Replies, one tab per category
  2) an interactive browse page (single .html file): search, category filters, sort, tap to see replies, star favorites
No usernames are included. Categories use keyword lists (no AI needed) — edit CATEGORIES below to change them.

Usage: python3 tools/comment_report.py comments.csv "Post title or link" [output_name]
"""
import csv, html, json, re, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

# ── categories: add/remove words freely. Matching is case-insensitive and on whole words (a * means "word starts with") ──
CATEGORIES = {
    "🏋️ Weight loss": ["lost \\d+", "lost (?:the |some |so much |all the )?weight", "\\d+ ?(?:lbs|lb|pounds|kg)", "weight ?loss", "lose weight", "losing weight", "overweight", "ozempic", "glp-?1", "wegovy", "mounjaro", "down \\d+ ?(?:lbs|pounds)?", "dress size", "pant size"],
    "🍭 Sugar & cravings": ["sugar", "sweets", "candy", "soda", "sodas", "craving*", "dessert*", "junk food", "binge*"],
    "🚶 Movement": ["walking", "started walking", "daily walks?", "go(?:ing)? for (?:a )?walks?", "gym", "pilates", "workouts?", "working out", "work out", "exercis\\w*", "yoga", "hik(?:e|es|ing)", "lifting", "weights", "5k", "10k", "marathon", "started running", "stairmaster", "treadmill", "\\d+k steps", "steps a day"],
    "🍷 Alcohol": ["sober", "sobriety", "stopped drinking", "quit drinking", "alcohol\\w*", "dry january", "booze", "\\baa\\b"],
    "🍽️ Eating habits": ["protein", "snack\\w*", "meal prep\\w*", "portions?", "diet", "dieting", "eating healthy", "healthy eating", "eat(?:ing)? clean", "clean eating", "cook(?:ing|ed)? (?:my|at home|meals)", "fast food", "vegan", "vegetarian", "fried food", "intermittent fasting", "calories"],
    "🧠 Mindset & mental health": ["anxiety", "depress\\w*", "therapy", "therapist", "confiden\\w*", "discipline", "mental health", "self ?worth", "self ?love", "healing", "psychiatrist", "counseling"],
    "💍 Relationships": ["husband", "wife", "married", "marry", "boyfriend", "girlfriend", "dating", "first date", "blind date", "my ex", "divorc\\w*", "fianc\\w*", "soulmate"],
    "💼 Career & money": ["job", "fired", "hired", "promot\\w*", "degree", "college", "grad school", "business", "career", "salary", "6 figures", "six figures", "boss", "quit my job", "interview"],
}
QUESTION = re.compile(r"\?\s*$|^(how|what|where|why|who|when|which|can|does|is|are|do|did)\b", re.I)
VIBES = re.compile(r"^(\W|lol|lmao|this|same|omg|me|yes|wow|haha\w*|😂|😭|🥹|❤️|🤣|\s)*$", re.I)

def compile_pat(words):
    parts = []
    for w in words:
        w = w.replace("*", "\\w*")
        parts.append(w)
    return re.compile(r"\b(" + "|".join(parts) + r")\b", re.I)
CAT_PATS = {k: compile_pat(v) for k, v in CATEGORIES.items()}

def categorize(text):
    cats = [k for k, p in CAT_PATS.items() if p.search(text)]
    if QUESTION.search(text.strip()):
        cats.append("❓ Questions")
    if not cats and (len(text.strip()) <= 12 or VIBES.match(text.strip())):
        cats.append("😂 Just vibes")
    return cats or ["Other"]

STOP = set("""a an the and or but so to of in on at for with from by is are was were be been am i me my we our you your he she it they them
this that these those just like get got have has had do did not no yes if then when than as up out about into over after before
one day all can will would could should really very too also im i'm it's its dont don't didn't cant can't what who how my
me him her his hers their there here more most some any every even still now ever never back went go going gone made make
said say told know knew think thought want wanted time year years lol""".split())

def top_phrases(texts, n=40):
    c = collections.Counter()
    for t in texts:
        words = [w for w in re.findall(r"[a-z0-9']+", t.lower())]
        for size in (2, 3):
            for i in range(len(words) - size + 1):
                g = words[i:i + size]
                if g[0] in STOP or g[-1] in STOP:
                    continue
                c[" ".join(g)] += 1
    return [(p, k) for p, k in c.most_common(n * 3) if k >= 3][:n]

def main():
    src, title = sys.argv[1], sys.argv[2]
    out = sys.argv[3] if len(sys.argv) > 3 else src.rsplit(".", 1)[0]
    rows, seen = [], set()
    for r in csv.DictReader(open(src, encoding="utf-8-sig")):  # drop duplicates TikTok sometimes repeats across pages
        if r["comment_id"] not in seen:
            seen.add(r["comment_id"]); rows.append(r)
    tops = sorted([r for r in rows if not r["reply_to"]], key=lambda r: -int(r["likes"]))
    num = {r["comment_id"]: i + 1 for i, r in enumerate(tops)}  # Comment # = rank by likes
    comments = [{"n": num[r["comment_id"]], "t": r["text"].strip(), "l": int(r["likes"]), "r": int(r["replies"] or 0),
                 "c": categorize(r["text"])} for r in tops]
    byn = {c["n"]: c for c in comments}
    replies = []
    for r in rows:
        if r["reply_to"] and r["reply_to"] in num:
            replies.append({"p": num[r["reply_to"]], "t": r["text"].strip(), "l": int(r["likes"])})
    replies.sort(key=lambda x: (x["p"], -x["l"]))
    cat_names = list(CATEGORIES) + ["❓ Questions", "😂 Just vibes", "Other"]
    counts = collections.Counter(c for x in comments for c in x["c"])
    phrases = top_phrases([c["t"] for c in comments])

    # ── Excel ──
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="C8325F"); wrap = Alignment(wrap_text=True, vertical="top")
    def sheet(ws, headers, data, widths):
        ws.append(headers)
        for i, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=i); cell.font = head; cell.fill = fill
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        for d in data:
            ws.append(d)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    ov = wb.active; ov.title = "Overview"
    ov.append(["TikTok comments report"]); ov["A1"].font = Font(bold=True, size=16, color="C8325F")
    ov.append(["Post", title]); ov.append(["Main comments", len(comments)]); ov.append(["Replies", len(replies)])
    ov.append([]); ov.append(["Category", "Comments"]); ov["A6"].font = Font(bold=True); ov["B6"].font = Font(bold=True)
    for k in cat_names:
        ov.append([k, counts.get(k, 0)])
    chart = BarChart(); chart.type = "bar"; chart.title = "Comments per category"; chart.legend = None
    chart.add_data(Reference(ov, min_col=2, min_row=6, max_row=6 + len(cat_names)), titles_from_data=True)
    chart.set_categories(Reference(ov, min_col=1, min_row=7, max_row=6 + len(cat_names)))
    chart.height, chart.width = 9, 16
    ov.add_chart(chart, "D2")
    start = 8 + len(cat_names)
    ov.cell(row=start, column=1, value="Phrases people keep saying").font = Font(bold=True)
    for i, (p, k) in enumerate(phrases):
        ov.cell(row=start + 1 + i, column=1, value=p); ov.cell(row=start + 1 + i, column=2, value=k)
    ov.column_dimensions["A"].width = 34; ov.column_dimensions["B"].width = 60

    cats_s = lambda c: ", ".join(c["c"])
    sheet(wb.create_sheet("Top Liked"), ["#", "Comment", "Likes", "Replies", "Category"],
          [[c["n"], c["t"], c["l"], c["r"], cats_s(c)] for c in comments[:100]], [6, 90, 10, 9, 30])
    sheet(wb.create_sheet("Comments"), ["Comment #", "Comment", "Likes", "# of replies", "Category"],
          [[c["n"], c["t"], c["l"], c["r"], cats_s(c)] for c in comments], [11, 90, 10, 12, 30])
    sheet(wb.create_sheet("Replies"), ["Reply to comment #", "Original comment (preview)", "Reply", "Likes"],
          [[x["p"], byn[x["p"]]["t"][:80] + ("…" if len(byn[x["p"]]["t"]) > 80 else ""), x["t"], x["l"]] for x in replies], [12, 45, 70, 10])
    for k in cat_names:
        items = [c for c in comments if k in c["c"]]
        if not items:
            continue
        name = re.sub(r"[^\w &-]", "", k).strip()[:30] or "Category"
        sheet(wb.create_sheet(name), ["Comment #", "Comment", "Likes", "# of replies"],
              [[c["n"], c["t"], c["l"], c["r"]] for c in items], [11, 95, 10, 12])
    wb.save(out + ".xlsx")

    # ── Interactive browse page ──
    data = {"title": title, "comments": comments, "replies": replies, "cats": cat_names,
            "counts": {k: counts.get(k, 0) for k in cat_names}, "phrases": phrases[:20]}
    page = PAGE.replace("__TITLE__", html.escape(title)).replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
    open(out + ".html", "w", encoding="utf-8").write(page)
    print(f"saved {out}.xlsx and {out}.html  ({len(comments)} comments, {len(replies)} replies)")
    for k in cat_names:
        print(f"  {k}: {counts.get(k, 0)}")

PAGE = r"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Comment Browser</title><style>
:root{--bg:#FBF5EF;--card:#fff;--ink:#2B1E23;--muted:#7B6A70;--line:#EBDCD6;--hot:#C8325F;--soft:#FCEDEA}
@media (prefers-color-scheme:dark){:root{--bg:#1c1518;--card:#271e22;--ink:#f3e9ec;--muted:#b19fa5;--line:#3a2d32;--soft:#352529}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:16px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}
header{position:sticky;top:0;z-index:5;background:var(--bg);padding:14px 16px 10px;border-bottom:1px solid var(--line)}
h1{font-size:18px;margin:0 0 2px}.sub{color:var(--muted);font-size:13px;margin-bottom:10px;word-break:break-all}
.bar{display:flex;gap:8px;flex-wrap:wrap}input[type=search]{flex:1;min-width:180px;padding:10px 12px;border:1px solid var(--line);border-radius:10px;background:var(--card);color:var(--ink);font-size:16px}
select{padding:10px;border:1px solid var(--line);border-radius:10px;background:var(--card);color:var(--ink);font-size:15px}
.chips{display:flex;gap:6px;overflow-x:auto;padding:10px 0 2px}.chip{flex:none;border:1px solid var(--line);background:var(--card);color:var(--ink);border-radius:99px;padding:6px 12px;font-size:14px;cursor:pointer}
.chip.on{background:var(--hot);border-color:var(--hot);color:#fff}.chip small{opacity:.7}
main{max-width:760px;margin:0 auto;padding:12px 16px 60px}.count{color:var(--muted);font-size:13px;margin:4px 0 10px}
.c{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:12px 14px;margin-bottom:10px}
.top{display:flex;gap:10px;align-items:flex-start}.n{flex:none;font-weight:700;color:var(--hot);font-size:13px;min-width:44px}
.t{flex:1;white-space:pre-wrap;word-wrap:break-word}.star{flex:none;border:0;background:none;font-size:20px;cursor:pointer;color:var(--muted);line-height:1}
.star.on{color:#E8A317}.meta{display:flex;gap:12px;flex-wrap:wrap;margin:8px 0 0 54px;font-size:13px;color:var(--muted)}
.tag{background:var(--soft);border-radius:99px;padding:1px 8px}.rbtn{border:0;background:none;color:var(--hot);font-weight:600;cursor:pointer;padding:0;font-size:13px}
.replies{margin:10px 0 0 54px;border-left:3px solid var(--line);padding-left:12px}.r{padding:6px 0;border-bottom:1px dashed var(--line);font-size:15px}
.r:last-child{border:0}.r small{color:var(--muted)}.more{display:block;margin:10px auto;padding:10px 18px;border:1px solid var(--line);border-radius:99px;background:var(--card);color:var(--ink);cursor:pointer}
details{margin:6px 0 0}summary{cursor:pointer;color:var(--muted);font-size:13px}.ph{display:flex;flex-wrap:wrap;gap:6px;margin-top:8px}.ph span{background:var(--soft);border-radius:99px;padding:2px 10px;font-size:13px;cursor:pointer}
mark{background:#FFE98F;color:inherit}
</style></head><body>
<header><h1>💬 Comment Browser</h1><div class="sub">__TITLE__</div>
<div class="bar"><input id="q" type="search" placeholder="Search comments (e.g. walking, lost 50)"><select id="sort"><option value="likes">Most liked</option><option value="replies">Most replies</option><option value="num">Comment #</option></select></div>
<div class="chips" id="chips"></div>
<details><summary>Phrases people keep saying</summary><div class="ph" id="ph"></div></details></header>
<main><div class="count" id="count"></div><div id="list"></div><button class="more" id="more">Show more</button></main>
<script>
const D=__DATA__;
const store={get(k){try{return JSON.parse(localStorage.getItem(k))}catch(e){return null}},set(k,v){try{localStorage.setItem(k,JSON.stringify(v))}catch(e){}}};
const KEY='stars:'+D.title; let stars=new Set(store.get(KEY)||[]);
const byP={};D.replies.forEach(r=>(byP[r.p]=byP[r.p]||[]).push(r));
let cat='All',q='',sort='likes',shown=60;const $=s=>document.querySelector(s);
const esc=s=>s.replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const hi=s=>{s=esc(s);if(!q)return s;return s.replace(new RegExp('('+q.replace(/[.*+?^${}()|[\]\\]/g,'\\$&')+')','gi'),'<mark>$1</mark>')};
function chips(){const all=[['All',D.comments.length],['⭐ Starred',stars.size],...D.cats.map(c=>[c,D.counts[c]])];
$('#chips').innerHTML=all.map(([c,n])=>`<button class="chip ${c===cat?'on':''}" data-c="${esc(c)}">${esc(c)} <small>${n}</small></button>`).join('');}
function list(){let xs=D.comments.filter(c=>(cat==='All'||(cat==='⭐ Starred'?stars.has(c.n):c.c.includes(cat)))&&(!q||c.t.toLowerCase().includes(q.toLowerCase())));
xs.sort(sort==='likes'?(a,b)=>b.l-a.l:sort==='replies'?(a,b)=>b.r-a.r:(a,b)=>a.n-b.n);
$('#count').textContent=xs.length.toLocaleString()+' comments';
$('#list').innerHTML=xs.slice(0,shown).map(c=>`<div class="c"><div class="top"><span class="n">#${c.n}</span><div class="t">${hi(c.t)}</div><button class="star ${stars.has(c.n)?'on':''}" data-s="${c.n}" title="Star">★</button></div>
<div class="meta"><span>❤️ ${c.l.toLocaleString()}</span>${c.c.map(t=>`<span class="tag">${esc(t)}</span>`).join('')}${byP[c.n]?`<button class="rbtn" data-r="${c.n}">💬 ${byP[c.n].length} replies</button>`:''}</div><div class="replies" id="r${c.n}" hidden></div></div>`).join('');
$('#more').hidden=xs.length<=shown;}
document.addEventListener('click',e=>{const t=e.target.closest('button,span[data-p]');if(!t)return;
if(t.dataset.c){cat=t.dataset.c;shown=60;chips();list();}
else if(t.dataset.s){const n=+t.dataset.s;stars.has(n)?stars.delete(n):stars.add(n);store.set(KEY,[...stars]);t.classList.toggle('on');chips();}
else if(t.dataset.r){const box=$('#r'+t.dataset.r);if(!box.innerHTML)box.innerHTML=byP[t.dataset.r].map(r=>`<div class="r">${esc(r.t)} <small>❤️ ${r.l}</small></div>`).join('');box.hidden=!box.hidden;}
else if(t.id==='more'){shown+=60;list();}
else if(t.dataset.p){$('#q').value=q=t.dataset.p;shown=60;list();}});
$('#q').addEventListener('input',e=>{q=e.target.value.trim();shown=60;list()});
$('#sort').addEventListener('change',e=>{sort=e.target.value;list()});
$('#ph').innerHTML=D.phrases.map(([p,k])=>`<span data-p="${esc(p)}">${esc(p)} · ${k}</span>`).join('');
chips();list();
</script></body></html>"""

if __name__ == "__main__":
    main()
